import io #lets us us in memory withot using a disk 
import pymupdf
import docx

def extract_pdf(file_bytes: bytes) -> list[dict]:
    """Read a PDF and return one item per page: {"page": 1, "text": "..."}."""
    pages = []
    with pymupdf.open(stream=file_bytes,filetype='pdf') as pdf:
        for page_number, page in enumerate(pdf,start=1):
            text = page.get_text('text',sort=True)
            pages.append({'page':page_number, 'text': text})
    return pages        


def extract_docx(file_bytes: bytes) -> list[dict]:
    """Read a DOCX and return its text. DOCX files have no fixed pages,
    so we return everything as one item with page = None."""
    document = docx.Document(io.BytesIO(file_bytes))
    lines = []
    for paragraph in document.paragraphs:
       lines.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    return [{"page": None, "text": "\n".join(lines)}]


def extract_text(file_bytes: bytes, extension: str) -> list[dict]:
    """Choose the right reader based on the file extension."""
    if extension == "pdf":
        return extract_pdf(file_bytes)
    if extension == "docx":
        return extract_docx(file_bytes)
    # We should never reach here, because main.py checks the extension first.
    raise ValueError(f"Unsupported file type: {extension}")
