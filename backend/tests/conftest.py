# conftest.py
# Shared test setup. pytest loads this file automatically before any test.
#
# Two important rules for our tests:
#   1. NO real Gemini calls: we replace the Gemini client with a FAKE one.
#   2. NO real data: ChromaDB and the active-resume file go to a temporary folder.

import io
import os
import socket
import tempfile
from types import SimpleNamespace

# These environment variables must be set BEFORE the app is imported,
# because config.py reads them at import time.
_TEMP_DIR = tempfile.mkdtemp(prefix="resume_tests_")
os.environ["CHROMA_DIR"] = os.path.join(_TEMP_DIR, "chroma")
os.environ["ACTIVE_RESUME_FILE"] = os.path.join(_TEMP_DIR, "active_resume.json")
os.environ["EVAL_RESULTS_DIR"] = os.path.join(_TEMP_DIR, "eval_results")
os.environ["GEMINI_API_KEY"] = "test-key-not-real"
os.environ["ANONYMIZED_TELEMETRY"] = "False"           # stop ChromaDB from sending telemetry

import docx                                            # noqa: E402  (import after env setup)
import pymupdf                                         # noqa: E402
import pytest                                          # noqa: E402
from fastapi.testclient import TestClient              # noqa: E402

from app import embeddings, logs, main, memory, rag    # noqa: E402

DIM = 64   # small fake embeddings are enough for tests


def fake_vector(text: str) -> list[float]:
    """A tiny "bag of words" embedding: each word adds 1 to one of 64 slots.
    Texts that share words get similar vectors, so search behaves sensibly."""
    vector = [0.0] * DIM
    for word in text.lower().replace(",", " ").replace(":", " ").split():
        vector[sum(map(ord, word)) % DIM] += 1.0
    vector[0] += 0.01                                  # never an all-zero vector
    return vector


class FakeModels:
    """Pretends to be client.models from the google-genai library."""

    def __init__(self):
        self.generate_calls = []                       # every prompt we were asked to answer
        self.fail_generate = False                     # set True to simulate "Gemini is down"

    def embed_content(self, model, contents, config):
        return SimpleNamespace(embeddings=[SimpleNamespace(values=fake_vector(t)) for t in contents])

    def generate_content(self, model, contents, config):
        self.generate_calls.append({"prompt": contents, "system": config.system_instruction})
        if self.fail_generate:
            raise RuntimeError("Gemini unavailable")
        if config.system_instruction == rag.REWRITE_PROMPT:
            return SimpleNamespace(text="Which programming languages were used in projects?")
        return SimpleNamespace(text="The candidate knows Python [1].")


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    """Safety net: any attempt to open an internet connection fails the test."""
    def blocked(*args, **kwargs):
        raise RuntimeError("Network access is not allowed in tests")
    monkeypatch.setattr(socket.socket, "connect", blocked)


@pytest.fixture
def fake_gemini(monkeypatch):
    """Replace the real Gemini client with the fake one."""
    fake = SimpleNamespace(models=FakeModels())
    monkeypatch.setattr(embeddings, "_client", fake)
    return fake.models


@pytest.fixture
def client(fake_gemini):
    """A test client with a clean app state for each test."""
    main.current_resume = None
    main.current_chunks = []
    memory.clear_all()
    logs.clear_logs()
    if os.path.exists(os.environ["ACTIVE_RESUME_FILE"]):
        os.remove(os.environ["ACTIVE_RESUME_FILE"])
    # raise_server_exceptions=False → crashes come back as HTTP 500, like a real server
    return TestClient(main.app, raise_server_exceptions=False)


# ---------------- Sample resume files, made in memory ----------------

RESUME_LINES = [
    "Jane Doe",
    "jane@example.com",
    "SKILLS",
    "Python, SQL, React, Docker",
    "PROJECTS",
    "Resume Chatbot",
    "Built a RAG chatbot in Python with FastAPI.",
    "Sales Dashboard",
    "Built a dashboard with React and SQL.",
    "EDUCATION",
    "BSc Physics, Pune University, 2020",
]


def make_pdf(lines=RESUME_LINES) -> bytes:
    """Create a small one-page PDF with the given lines."""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "\n".join(lines))
    return document.tobytes()


def make_docx(lines=RESUME_LINES) -> bytes:
    """Create a small DOCX: paragraphs plus one table (tables must be read too)."""
    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Languages: English, Hindi"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def uploaded(client):
    """A client that already has the sample PDF resume indexed."""
    response = client.post("/resume", files={"file": ("resume.pdf", make_pdf(), "application/pdf")})
    assert response.status_code == 201
    return client
